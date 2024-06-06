import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document
from pathlib import Path
from pandas import read_csv
from typing import Dict

def fill_template(template_path: Path, output_path: Path, data: Dict[str, str]) -> None:
    """Fill the Word document template with provided data and save it to the output path."""
    doc = Document(template_path)
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            if key in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(key, value)

    doc.save(output_path)

def generate_docs_from_csv(csv_file: Path, template_path: Path, output_dir: Path) -> None:
    """Generate Word documents for each row in the CSV file using the template."""
    df = read_csv(csv_file)
    output_dir.mkdir(exist_ok=True)

    for index, row in df.iterrows():
        data = {f'[{col}]': str(row[col]) for col in df.columns}
        recipient_name = row.get('Recipient', f"Invitation_{index + 1}")
        output_path = output_dir / f"{recipient_name}_invitation.docx"
        fill_template(template_path=template_path, output_path=output_path, data=data)

def main(csv_file_path: Path, template_file_path: Path, output_dir_path: Path) -> None:
    """Main function to handle the document generation process."""
    if not csv_file_path.exists():
        raise FileNotFoundError(f"CSV file not found: {csv_file_path}")
    if not template_file_path.exists():
        raise FileNotFoundError(f"Template file not found: {template_file_path}")

    generate_docs_from_csv(csv_file=csv_file_path, template_path=template_file_path, output_dir=output_dir_path)

def select_csv_file():
    csv_file_path.set(filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")]))

def select_template_file():
    template_file_path.set(filedialog.askopenfilename(filetypes=[("Word files", "*.docx")]))

def select_output_dir():
    output_dir_path.set(filedialog.askdirectory())

def run_generation():
    try:
        csv_path = Path(csv_file_path.get())
        template_path = Path(template_file_path.get())
        output_path = Path(output_dir_path.get())
        main(csv_file_path=csv_path, template_file_path=template_path, output_dir_path=output_path)
        messagebox.showinfo("Success", "Documents generated successfully!")
    except Exception as e:
        messagebox.showerror("Error", str(e))

def show_help():
    help_text = (
        "How to use the Document Generator:\n\n"
        "1. Select the CSV file containing the data:\n"
        "   - Click 'Browse' next to 'Select CSV File' and choose your CSV file.\n\n"
        "2. Select the Word template file with placeholders:\n"
        "   - Click 'Browse' next to 'Select Template File' and choose your template file.\n\n"
        "3. Select the output directory where documents will be saved:\n"
        "   - Click 'Browse' next to 'Select Output Directory' and choose the directory.\n\n"
        "4. Click 'Generate Documents' to create the documents based on your CSV and template.\n\n"
        "If a column in the CSV does not match any placeholder in the template, it will be ignored."
    )
    messagebox.showinfo("Help", help_text)

def open_link(event):
    import webbrowser
    webbrowser.open_new(event.widget.cget("text"))


def show_credits():
    credits_win = tk.Toplevel(root)
    credits_win.title("Credits")
    
    tk.Label(credits_win, text="This Document Generator was created by @PawanPerera.\n\nFor more information, contact: 1pawanpinsara@gmail.com", padx=20, pady=20).pack()
    link = tk.Label(credits_win, text="https://github.com/PinsaraPerera", fg="blue", cursor="hand2")
    link.pack(padx=20, pady=(0, 20))
    link.bind("<Button-1>", open_link)


def exit():
    root.destroy()

# Set up the GUI
root = tk.Tk()
root.title("Document Generator")

csv_file_path = tk.StringVar()
template_file_path = tk.StringVar()
output_dir_path = tk.StringVar()

tk.Label(root, text="Select CSV File:").grid(row=0, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=csv_file_path, width=50).grid(row=0, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_csv_file).grid(row=0, column=2, padx=10, pady=10)

tk.Label(root, text="Select Template File:").grid(row=1, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=template_file_path, width=50).grid(row=1, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_template_file).grid(row=1, column=2, padx=10, pady=10)

tk.Label(root, text="Select Output Directory:").grid(row=2, column=0, padx=10, pady=10)
tk.Entry(root, textvariable=output_dir_path, width=50).grid(row=2, column=1, padx=10, pady=10)
tk.Button(root, text="Browse", command=select_output_dir).grid(row=2, column=2, padx=10, pady=10)

tk.Button(root, text="Generate Documents", command=run_generation).grid(row=3, column=1, padx=10, pady=20)

# Create a menubar
menubar = tk.Menu(root)

# Add File menu with Exit option
file_menu = tk.Menu(menubar, tearoff=0)
file_menu.add_command(label="Exit", command=exit)
menubar.add_cascade(label="File", menu=file_menu)

# Add a Help menu
help_menu = tk.Menu(menubar, tearoff=0)
help_menu.add_command(label="Help", command=show_help)
menubar.add_cascade(label="Help", menu=help_menu)

# Add Credit menu
credit_menu = tk.Menu(menubar, tearoff=0)
credit_menu.add_command(label="Credits", command=show_credits)
menubar.add_cascade(label="Credits", menu=credit_menu)

# Add credits to the Help menu
help_menu.add_separator()
help_menu.add_command(label="Credits", command=show_credits)


# Highlight the Help menu
menubar.entryconfig("Help", font=("Helvetica", 10, "bold"), background="yellow")

root.config(menu=menubar)
root.mainloop()
